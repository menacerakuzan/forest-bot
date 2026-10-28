import io
import shutil
from pathlib import Path

import docx
import pytest

from forest_bot.extract import (
    EmptyDocumentError,
    UnsupportedFormatError,
    _clean_ocr,
    extract_text,
    resolve_suffix,
    split_into_chunks,
)

ARCHIVE = Path(__file__).parent.parent / "data" / "archive"


def test_txt_utf8_and_cp1251() -> None:
    text = "Внесок — 4 жолуді"
    assert extract_text(text.encode(), ".txt") == text
    assert extract_text(text.encode("cp1251"), ".txt") == text


def test_docx_paragraphs_and_tables() -> None:
    document = docx.Document()
    document.add_paragraph("Правила")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Тиша"
    table.rows[0].cells[1].text = "вночі"
    buffer = io.BytesIO()
    document.save(buffer)
    assert extract_text(buffer.getvalue(), ".docx") == "Правила\nТиша | вночі"


def test_pdf_with_text_layer() -> None:
    text = extract_text((ARCHIVE / "Карта лісу.pdf").read_bytes(), ".pdf")
    assert "Під вербою — і листяний ліс, і біля річки" in text


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract is not installed")
def test_scanned_pdf_is_ocred() -> None:
    pdf = ARCHIVE / "Членські внески рукавички Під ялиною.pdf"
    text = extract_text(pdf.read_bytes(), ".pdf")
    assert "3 шишки" in text
    assert "Білочці" in text


def test_unsupported_and_empty() -> None:
    with pytest.raises(UnsupportedFormatError):
        extract_text(b"x", ".doc")
    with pytest.raises(EmptyDocumentError):
        extract_text(b"  \n ", ".txt")


def test_resolve_suffix_falls_back_to_mime() -> None:
    assert resolve_suffix("Порада.PDF", None) == ".pdf"
    assert resolve_suffix(None, "text/plain") == ".txt"


def test_split_into_chunks_respects_limit() -> None:
    text = "\n".join(f"абзац {i} " + "слово " * 20 for i in range(50))
    chunks = split_into_chunks(text, max_chars=500)
    assert len(chunks) > 1
    assert all(len(c) <= 500 for c in chunks)
    assert "".join(chunks).replace("\n", "") == text.replace("\n", "")


def test_clean_ocr_drops_noise_lines() -> None:
    raw = "— пот - —_ Й\nСкільки і коли\n~ —\nВнесок — 3 шишки. _-4\nри . ---"
    assert _clean_ocr(raw) == "Скільки і коли\nВнесок — 3 шишки. _-4"
