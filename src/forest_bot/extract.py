"""Витягування тексту з PDF, Word і текстових файлів."""

import io
import logging
import re
from pathlib import Path

import docx
from pypdf import PdfReader

log = logging.getLogger(__name__)

# Менше тексту в PDF — вважаємо його сканом і запускаємо OCR.
MIN_TEXT_LAYER_CHARS = 50
OCR_DPI = 400
OCR_THRESHOLD = 150
OCR_LANGUAGES = "ukr"

SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".txt", ".md"})

MIME_SUFFIXES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
    "text/markdown": ".md",
}


class UnsupportedFormatError(Exception):
    pass


class EmptyDocumentError(Exception):
    pass


def resolve_suffix(file_name: str | None, mime_type: str | None) -> str:
    suffix = Path(file_name or "").suffix.lower()
    if not suffix and mime_type:
        suffix = MIME_SUFFIXES.get(mime_type, "")
    return suffix


def extract_text(data: bytes, suffix: str) -> str:
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormatError(suffix)
    if suffix == ".pdf":
        raw = _pdf(data)
    elif suffix == ".docx":
        raw = _docx(data)
    else:
        raw = _plain(data)
    text = normalize(raw)
    if not text:
        raise EmptyDocumentError
    return text


def _pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if len(normalize(text)) >= MIN_TEXT_LAYER_CHARS:
        return text
    # Тексту майже немає — це скан, розпізнаємо сторінки.
    return _ocr_pdf(data)


def _ocr_pdf(data: bytes) -> str:
    import pypdfium2
    import pytesseract

    pdf = pypdfium2.PdfDocument(data)
    try:
        pages = []
        for page in pdf:
            image = page.render(scale=OCR_DPI / 72).to_pil().convert("L")
            # Бінаризація прибирає світлий фон (папір, текстури), лишаючи темний текст.
            image = image.point(lambda p: 255 if p > OCR_THRESHOLD else 0)
            raw = pytesseract.image_to_string(image, lang=OCR_LANGUAGES, config="--psm 6")
            pages.append(_clean_ocr(raw))
        return "\n\n".join(pages)
    except pytesseract.TesseractNotFoundError:
        log.error("tesseract is not installed — scanned PDFs cannot be read")
        return ""
    finally:
        pdf.close()


_OCR_EDGE_JUNK = " \t—–-_~|=:;\"'“”„.,*`^"
_WORD = re.compile(r"[^\W\d_]{4,}")


def _clean_ocr(text: str) -> str:
    """Прибирає рядки-шум, які OCR «бачить» у візерунках фону."""
    lines = []
    for line in text.splitlines():
        line = line.strip(_OCR_EDGE_JUNK)
        meaningful = sum(ch.isalnum() for ch in line)
        if line and _WORD.search(line) and meaningful / len(line) >= 0.5:
            lines.append(line)
    return "\n".join(lines)


def _docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _plain(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_into_chunks(text: str, max_chars: int = 1500) -> list[str]:
    """Ріже великий документ на шматки по абзацах, щоб не перевищувати ліміти LLM."""
    pieces = [
        paragraph[i : i + max_chars]
        for paragraph in text.split("\n")
        for i in range(0, max(len(paragraph), 1), max_chars)
    ]
    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + len(piece) + 1 > max_chars:
            chunks.append(current)
            current = piece
        else:
            current = f"{current}\n{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks
